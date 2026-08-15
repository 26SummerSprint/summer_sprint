"""
Judge LLM 성능 평가 스크립트.

입력:
    1. profiles.json
    2. P1.csv ~ P12.csv 또는 P1.xlsx ~ P12.xlsx

Gold Set 형식:

    arxiv_id,title,abstract_clean,label

예:
    2401.00001,Paper A,Abstract A...,1
    2401.00002,Paper B,Abstract B...,0

label:
    1 = 사람 기준 정답 논문
    0 = 사람 기준 오답 논문

Judge에는 human label을 전달하지 않는다.

실행:

전체 평가:

    python -m ground_truth_labeler.evaluate_judge ^
        --input-dir ./gold_sets ^
        --profile-file ./profiles.json ^
        --output-dir ./judge_eval

특정 프로필:

    python -m ground_truth_labeler.evaluate_judge ^
        --input ./gold_sets/P1.xlsx ^
        --profile-file ./profiles.json ^
        --profile P1 ^
        --output-dir ./judge_eval

두 단계로 분리해서 실행 (main.py로 먼저 논문 평가, 그 결과로 성능평가만):

    1) 논문 평가 (main.py, Gemini 호출):

        python -m ground_truth_labeler.main ^
            --profiles-file ./profiles.json ^
            --input-dir ./gold_sets ^
            --output-dir ./judge_labels ^
            --keep-raw

    2) 성능평가 (evaluate_judge.py, Gemini 호출 없이 1)의 결과만 사용):

        python -m ground_truth_labeler.evaluate_judge ^
            --input-dir ./gold_sets ^
            --profile-file ./profiles.json ^
            --labels-dir ./judge_labels ^
            --output-dir ./judge_eval

지원 파일:
    .csv
    .xlsx
"""

import argparse
import asyncio
import csv
import json
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import load_workbook
from dataclasses import asdict, is_dataclass
from .config import GEMINI_API_KEY
from .labeler import GroundTruthLabeler


# ============================================================
# 필수 컬럼
# ============================================================

REQUIRED_COLUMNS = {
    "arxiv_id",
    "title",
    "abstract_clean",
    "label",
}


# ============================================================
# 프로필 로딩
# ============================================================

def load_profiles(path: Path) -> List[Dict[str, Any]]:
    """
    profiles.json을 읽는다.

    예상 구조:

    [
        {
            "profile_id": "P1",
            "profile_text": "...",
            "keywords": [...],
            "exclusion": "..."
        },
        ...
    ]
    """

    if not path.exists():
        raise FileNotFoundError(
            f"프로필 파일을 찾을 수 없습니다: {path}"
        )

    try:
        with path.open(
            "r",
            encoding="utf-8-sig",
        ) as f:
            data = json.load(f)

    except json.JSONDecodeError as e:
        raise ValueError(
            f"프로필 JSON 형식이 잘못되었습니다: {path}\n"
            f"{e}"
        ) from e

    if not isinstance(data, list):
        raise ValueError(
            "profiles.json의 최상위 구조는 JSON 배열이어야 합니다.\n"
            "예: [{...}, {...}]"
        )

    profiles = []

    for index, profile in enumerate(data, start=1):

        if not isinstance(profile, dict):
            raise ValueError(
                f"profiles.json의 {index}번째 항목이 객체가 아닙니다."
            )

        profile_id = str(
            profile.get("profile_id", "")
        ).strip()

        if not profile_id:
            raise ValueError(
                f"profiles.json의 {index}번째 프로필에 "
                f"profile_id가 없습니다."
            )

        profile_text = str(
            profile.get("profile_text", "")
        ).strip()

        if not profile_text:
            raise ValueError(
                f"{profile_id}의 profile_text가 비어 있습니다."
            )

        profiles.append(profile)

    return profiles


# ============================================================
# 컬럼 검증
# ============================================================

def validate_columns(
    columns: List[str],
    path: Path,
) -> None:
    """
    Gold Set에 필요한 컬럼이 존재하는지 검사한다.
    """

    normalized = [
        str(column).strip()
        for column in columns
        if column is not None
    ]

    column_set = set(normalized)

    missing = REQUIRED_COLUMNS - column_set

    if missing:
        raise ValueError(
            f"{path.name}에 필요한 컬럼이 없습니다.\n"
            f"누락된 컬럼: {', '.join(sorted(missing))}\n"
            f"필요한 컬럼: {', '.join(sorted(REQUIRED_COLUMNS))}"
        )


# ============================================================
# Gold row 파싱
# ============================================================

def parse_gold_row(
    row: Dict[str, Any],
    path: Path,
    row_num: int,
) -> Dict[str, Any]:
    """
    CSV/XLSX 한 행을 공통 형식으로 변환한다.
    """

    arxiv_id = str(
        row.get("arxiv_id", "") or ""
    ).strip()

    title = str(
        row.get("title", "") or ""
    ).strip()

    abstract = str(
        row.get("abstract_clean", "") or ""
    ).strip()

    if not arxiv_id:
        raise ValueError(
            f"{path.name}:{row_num} "
            f"arxiv_id가 비어 있습니다."
        )

    if not title:
        raise ValueError(
            f"{path.name}:{row_num} "
            f"title이 비어 있습니다."
        )

    raw_label = row.get("label")

    try:

        # Excel에서 1.0 / 0.0 형태로 들어오는 경우 처리
        if isinstance(raw_label, float):
            if raw_label.is_integer():
                raw_label = int(raw_label)

        label = int(
            str(raw_label).strip()
        )

    except (
        ValueError,
        TypeError,
    ) as e:

        raise ValueError(
            f"{path.name}:{row_num}의 label이 "
            f"0 또는 1이 아닙니다: {raw_label}"
        ) from e

    if label not in (0, 1):
        raise ValueError(
            f"{path.name}:{row_num}의 label은 "
            f"0 또는 1이어야 합니다."
        )

    return {
        "arxiv_id": arxiv_id,
        "title": title,
        "abstract_clean": abstract,
        "human_label": label,
    }


# ============================================================
# CSV 로딩
# ============================================================

def load_gold_csv(
    path: Path,
) -> List[Dict[str, Any]]:
    """
    CSV Gold Set을 읽는다.
    """

    if not path.exists():
        raise FileNotFoundError(
            f"파일을 찾을 수 없습니다: {path}"
        )

    with path.open(
        "r",
        encoding="utf-8-sig",
        newline="",
    ) as f:

        reader = csv.DictReader(f)

        if reader.fieldnames is None:
            raise ValueError(
                f"CSV 헤더가 없습니다: {path}"
            )

        validate_columns(
            reader.fieldnames,
            path,
        )

        rows = []

        for row_num, row in enumerate(
            reader,
            start=2,
        ):

            rows.append(
                parse_gold_row(
                    row,
                    path,
                    row_num,
                )
            )

    return rows


# ============================================================
# XLSX 로딩
# ============================================================

def load_gold_xlsx(
    path: Path,
) -> List[Dict[str, Any]]:
    """
    XLSX Gold Set을 읽는다.

    첫 번째 worksheet를 사용한다.
    """

    if not path.exists():
        raise FileNotFoundError(
            f"파일을 찾을 수 없습니다: {path}"
        )

    workbook = load_workbook(
        filename=path,
        read_only=True,
        data_only=True,
    )

    try:

        worksheet = workbook.active

        rows_iter = worksheet.iter_rows(
            values_only=True
        )

        try:
            header = next(rows_iter)

        except StopIteration:
            raise ValueError(
                f"Excel 파일이 비어 있습니다: {path}"
            )

        columns = [
            str(value).strip()
            if value is not None
            else ""
            for value in header
        ]

        validate_columns(
            columns,
            path,
        )

        rows = []

        for row_num, values in enumerate(
            rows_iter,
            start=2,
        ):

            row = {
                columns[i]: values[i]
                for i in range(
                    min(
                        len(columns),
                        len(values),
                    )
                )
            }

            # 완전히 빈 행은 무시
            if not any(
                value is not None
                and str(value).strip() != ""
                for value in values
            ):
                continue

            rows.append(
                parse_gold_row(
                    row,
                    path,
                    row_num,
                )
            )

        return rows

    finally:
        workbook.close()


# ============================================================
# 파일 자동 로딩
# ============================================================

def load_gold_file(
    path: Path,
) -> List[Dict[str, Any]]:
    """
    확장자에 따라 CSV/XLSX를 자동으로 읽는다.
    """

    suffix = path.suffix.lower()

    if suffix == ".csv":
        return load_gold_csv(path)

    if suffix == ".xlsx":
        return load_gold_xlsx(path)

    raise ValueError(
        f"지원하지 않는 파일 형식입니다: {path}\n"
        f"지원 형식: .csv, .xlsx"
    )


# ============================================================
# Judge 실행
# ============================================================

async def run_judge(
    labeler: GroundTruthLabeler,
    profile: str,
    papers: List[Dict[str, Any]],
) -> List[Dict[str, Any]]:
    """
    Judge에 논문을 전달한다.

    중요:
        human_label은 절대로 Judge 입력에 포함하지 않는다.
    """

    judge_papers = [
        {
            "arxiv_id": row["arxiv_id"],
            "title": row["title"],
            "abstract_clean": row["abstract_clean"],
        }
        for row in papers
    ]

    def on_progress(
        completed: int,
        total: int,
        title: str,
        success: bool,
    ) -> None:

        status = (
            "완료"
            if success
            else "실패"
        )

        print(
            f"[Judge {completed}/{total}] "
            f"{title} ... {status}",
            file=sys.stderr,
        )

    raw_results = await labeler.label_batch(
        profile,
        judge_papers,
        title_key="title",
        abstract_key="abstract_clean",
        id_key="arxiv_id",
        on_progress=on_progress,
    )

    return raw_results


# ============================================================
# Judge label 추출
# ============================================================

def extract_judge_label(
    result: Optional[Any],
) -> Optional[int]:
    """
    Judge 결과에서 label을 추출한다.

    LabelResult 객체와 dict 모두 지원한다.
    """

    if not result:
        return None

    # LabelResult 객체
    if hasattr(result, "label"):
        value = result.label

    # dict
    elif isinstance(result, dict):
        value = result.get("label")

    else:
        return None

    try:
        value = int(value)

    except (
        TypeError,
        ValueError,
    ):
        return None

    if value not in (0, 1):
        return None

    return value


# ============================================================
# Human / Judge 결과 병합
# ============================================================

def merge_results(
    gold_rows: List[Dict[str, Any]],
    judge_results: List[Dict[str, Any]],
    profile_name: str,
) -> List[Dict[str, Any]]:
    """
    arxiv_id 기준으로 Human Label과 Judge Label을 합친다.
    """

    gold_by_id = {
        row["arxiv_id"]: row
        for row in gold_rows
    }

    merged = []

    for result in judge_results:

        arxiv_id = result.get("arxiv_id")

        gold = gold_by_id.get(arxiv_id)

        if gold is None:
            continue

        judge_result = result.get("result")

        judge_label = extract_judge_label(
            judge_result
        )

        # LabelResult 객체
        if judge_result is not None and hasattr(
            judge_result,
            "parse_warnings",
        ):
            parse_warnings = (
                judge_result.parse_warnings
            )

        # 혹시 dict가 들어오는 경우
        elif isinstance(judge_result, dict):
            parse_warnings = (
                judge_result.get(
                    "parse_warnings",
                    [],
                )
            )

        else:
            parse_warnings = []

        merged.append(
            {
                "profile": profile_name,
                "arxiv_id": arxiv_id,
                "title": gold["title"],
                "human_label": gold["human_label"],
                "judge_label": judge_label,
                "error": result.get("error"),
                "parse_warnings": parse_warnings,
            }
        )

    return merged


# ============================================================
# Confusion Matrix
# ============================================================

def confusion_matrix(
    rows: List[Dict[str, Any]],
) -> Dict[str, int]:

    tp = 0
    tn = 0
    fp = 0
    fn = 0

    for row in rows:

        human = row["human_label"]
        judge = row["judge_label"]

        if judge is None:
            continue

        if human == 1 and judge == 1:
            tp += 1

        elif human == 0 and judge == 0:
            tn += 1

        elif human == 0 and judge == 1:
            fp += 1

        elif human == 1 and judge == 0:
            fn += 1

    return {
        "TP": tp,
        "TN": tn,
        "FP": fp,
        "FN": fn,
    }


# ============================================================
# Metrics
# ============================================================

def calculate_metrics(
    rows: List[Dict[str, Any]],
) -> Dict[str, Any]:

    cm = confusion_matrix(
        rows
    )

    tp = cm["TP"]
    tn = cm["TN"]
    fp = cm["FP"]
    fn = cm["FN"]

    evaluated = (
        tp
        + tn
        + fp
        + fn
    )

    precision = (
        tp / (tp + fp)
        if (tp + fp) > 0
        else 0.0
    )

    recall = (
        tp / (tp + fn)
        if (tp + fn) > 0
        else 0.0
    )

    f1 = (
        2
        * precision
        * recall
        / (precision + recall)
        if (precision + recall) > 0
        else 0.0
    )

    accuracy = (
        (tp + tn) / evaluated
        if evaluated > 0
        else 0.0
    )

    positive_count = sum(
        1
        for row in rows
        if row["human_label"] == 1
    )

    judge_positive_count = sum(
        1
        for row in rows
        if row["judge_label"] == 1
    )

    disagreement_count = sum(
        1
        for row in rows
        if (
            row["judge_label"] is not None
            and row["human_label"]
            != row["judge_label"]
        )
    )

    failed_count = sum(
        1
        for row in rows
        if row.get("error")
    )

    return {
        "total": len(rows),
        "evaluated": evaluated,
        "failed": failed_count,
        "positive": positive_count,
        "judge_positive": judge_positive_count,
        "disagreement": disagreement_count,
        "accuracy": accuracy,
        "precision": precision,
        "recall": recall,
        "f1": f1,
        **cm,
    }


# ============================================================
# Per Profile
# ============================================================

def calculate_per_profile(
    rows: List[Dict[str, Any]],
) -> Dict[str, Dict[str, Any]]:

    profiles = sorted(
        set(
            row["profile"]
            for row in rows
        )
    )

    output = {}

    for profile in profiles:

        profile_rows = [
            row
            for row in rows
            if row["profile"] == profile
        ]

        output[profile] = calculate_metrics(
            profile_rows
        )

    return output


# ============================================================
# 불일치 추출
# ============================================================

def get_disagreements(
    rows: List[Dict[str, Any]],
) -> List[Dict[str, Any]]:

    return [
        row
        for row in rows
        if (
            row["judge_label"] is not None
            and row["human_label"]
            != row["judge_label"]
        )
    ]


# ============================================================
# CSV 저장
# ============================================================

def write_csv(
    rows: List[Dict[str, Any]],
    path: Path,
) -> None:

    if not rows:
        return

    fieldnames = list(
        rows[0].keys()
    )

    with path.open(
        "w",
        encoding="utf-8-sig",
        newline="",
    ) as f:

        writer = csv.DictWriter(
            f,
            fieldnames=fieldnames,
        )

        writer.writeheader()

        for row in rows:
            writer.writerow(row)


# ============================================================
# JSON 저장
# ============================================================

def write_json(
    data: Any,
    path: Path,
) -> None:

    def convert(obj):
        if is_dataclass(obj):
            return asdict(obj)

        if isinstance(obj, dict):
            return {
                key: convert(value)
                for key, value in obj.items()
            }

        if isinstance(obj, list):
            return [
                convert(value)
                for value in obj
            ]

        return obj

    path.write_text(
        json.dumps(
            convert(data),
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )


# ============================================================
# Metrics 출력
# ============================================================

def print_metrics(
    title: str,
    metrics: Dict[str, Any],
) -> None:

    print()
    print("=" * 60)
    print(title)
    print("=" * 60)

    print(
        f"전체             : {metrics['total']}"
    )

    print(
        f"평가 가능         : {metrics['evaluated']}"
    )

    print(
        f"실패             : {metrics['failed']}"
    )

    print(
        f"Human Positive    : {metrics['positive']}"
    )

    print(
        f"Judge Positive    : {metrics['judge_positive']}"
    )

    print(
        f"불일치            : {metrics['disagreement']}"
    )

    print()

    print(
        f"Accuracy          : "
        f"{metrics['accuracy']:.4f}"
    )

    print(
        f"Precision         : "
        f"{metrics['precision']:.4f}"
    )

    print(
        f"Recall            : "
        f"{metrics['recall']:.4f}"
    )

    print(
        f"F1                : "
        f"{metrics['f1']:.4f}"
    )

    print()

    print(
        f"TP={metrics['TP']} "
        f"TN={metrics['TN']} "
        f"FP={metrics['FP']} "
        f"FN={metrics['FN']}"
    )


# ============================================================
# Per Profile 출력
# ============================================================

def print_per_profile(
    per_profile: Dict[str, Dict[str, Any]],
) -> None:

    print()
    print("=" * 100)
    print("Per Profile")
    print("=" * 100)

    print(
        f"{'Profile':<10}"
        f"{'N':>7}"
        f"{'Pos':>7}"
        f"{'Precision':>12}"
        f"{'Recall':>10}"
        f"{'F1':>10}"
        f"{'Disagree':>10}"
    )

    print("-" * 100)

    for profile, metrics in per_profile.items():

        print(
            f"{profile:<10}"
            f"{metrics['evaluated']:>7}"
            f"{metrics['positive']:>7}"
            f"{metrics['precision']:>12.4f}"
            f"{metrics['recall']:>10.4f}"
            f"{metrics['f1']:>10.4f}"
            f"{metrics['disagreement']:>10}"
        )


# ============================================================
# main.py가 미리 저장해둔 라벨링 결과 로딩
# ============================================================

def load_precomputed_judge_results(
    path: Path,
) -> List[Dict[str, Any]]:
    """
    main.py --output(또는 --output-dir 다수 프로필 모드)이 저장한
    JSON 결과 파일을 읽는다.

    형식 (main.py write_output()의 결과):

        [
            {
                "arxiv_id": "...",
                "title": "...",
                "result": {"score": ..., "decision": ..., "label": ..., ...},
                "error": null
            },
            ...
        ]

    merge_results()/extract_judge_label()이 dict형 result도 그대로
    처리하므로 별도 변환 없이 사용할 수 있다.
    """

    if not path.exists():
        raise FileNotFoundError(
            f"라벨링 결과 파일을 찾을 수 없습니다: {path}\n"
            f"(main.py로 먼저 논문 평가를 실행해 이 파일을 만들어야 합니다.)"
        )

    with path.open("r", encoding="utf-8-sig") as f:
        data = json.load(f)

    if not isinstance(data, list):
        raise ValueError(
            f"{path.name}은 JSON 배열이어야 합니다 "
            f"(main.py --output 결과 형식)."
        )

    return data


# ============================================================
# 프로필 하나 평가
# ============================================================

async def evaluate_profile(
    labeler: Optional[GroundTruthLabeler],
    input_path: Path,
    profile: str,
    profile_name: str,
    output_dir: Path,
    labels_dir: Optional[Path] = None,
) -> List[Dict[str, Any]]:

    print()
    print(
        "=" * 60,
        file=sys.stderr,
    )

    print(
        f"{profile_name} Judge 평가 시작",
        file=sys.stderr,
    )

    print(
        "=" * 60,
        file=sys.stderr,
    )

    # CSV 또는 XLSX 자동 로딩 (사람 label은 여기서만 사용, Judge에는 전달 안 함)
    gold_rows = load_gold_file(
        input_path
    )

    print(
        f"입력 파일: {input_path}",
        file=sys.stderr,
    )

    print(
        f"Gold Set: {len(gold_rows)}건",
        file=sys.stderr,
    )

    # --------------------------------------------------------
    # Judge 결과: 이미 라벨링된 파일이 있으면 그것을 쓰고(Gemini 호출 없음),
    # 없으면 지금 바로 Gemini를 호출한다.
    # --------------------------------------------------------

    if labels_dir is not None:

        labels_path = labels_dir / f"{profile_name}.json"

        print(
            f"[{profile_name}] 사전 계산된 라벨링 결과 사용: "
            f"{labels_path}",
            file=sys.stderr,
        )

        judge_results = load_precomputed_judge_results(
            labels_path
        )

    else:

        judge_results = await run_judge(
            labeler,
            profile,
            gold_rows,
        )

    # --------------------------------------------------------
    # Human vs Judge 병합
    # --------------------------------------------------------

    merged = merge_results(
        gold_rows,
        judge_results,
        profile_name,
    )

    # --------------------------------------------------------
    # 전체 비교 결과
    # --------------------------------------------------------

    write_csv(
        merged,
        output_dir
        / f"{profile_name}_comparison.csv",
    )

    # --------------------------------------------------------
    # 불일치 결과
    # --------------------------------------------------------

    disagreements = get_disagreements(
        merged
    )

    write_csv(
        disagreements,
        output_dir
        / f"{profile_name}_disagreements.csv",
    )

    # --------------------------------------------------------
    # Judge 원본 결과
    # --------------------------------------------------------

    write_json(
        judge_results,
        output_dir
        / f"{profile_name}_judge_results.json",
    )

    # --------------------------------------------------------
    # Metrics
    # --------------------------------------------------------

    metrics = calculate_metrics(
        merged
    )

    write_json(
        metrics,
        output_dir
        / f"{profile_name}_metrics.json",
    )

    print_metrics(
        profile_name,
        metrics,
    )

    return merged


# ============================================================
# Argument Parser
# ============================================================

def build_parser() -> argparse.ArgumentParser:

    parser = argparse.ArgumentParser(
        description=(
            "P1~P12 사람 Gold Set과 "
            "Gemini Judge 결과를 자동 비교한다. "
            "CSV와 XLSX를 모두 지원한다."
        )
    )

    parser.add_argument(
        "--input-dir",
        default="./gold_sets",
        help=(
            "P1.csv~P12.csv 또는 "
            "P1.xlsx~P12.xlsx가 있는 디렉터리"
        ),
    )

    parser.add_argument(
        "--profile-file",
        required=True,
        help=(
            "P1~P12 프로필이 들어 있는 "
            "profiles.json"
        ),
    )

    parser.add_argument(
        "--input",
        help=(
            "특정 CSV/XLSX 하나만 평가할 때 사용"
        ),
    )

    parser.add_argument(
        "--profile",
        help=(
            "특정 프로필 하나만 평가. "
            "예: P1"
        ),
    )

    parser.add_argument(
        "--output-dir",
        default="./judge_eval",
        help=(
            "평가 결과 저장 디렉터리"
        ),
    )

    parser.add_argument(
        "--labels-dir",
        help=(
            "main.py가 미리 만들어둔 라벨링 결과 디렉터리 "
            "({labels-dir}/{profile_id}.json, main.py --output-dir와 동일 형식). "
            "지정하면 Gemini를 다시 호출하지 않고 이 파일들로 "
            "성능평가(사람 label과 비교)만 수행한다. "
            "지정하지 않으면 기존처럼 여기서 직접 Gemini를 호출한다."
        ),
    )

    return parser


# ============================================================
# Main
# ============================================================

async def async_main() -> None:

    parser = build_parser()
    args = parser.parse_args()

    # --------------------------------------------------------
    # labels-dir 모드가 아닐 때만 Gemini API Key가 필요하다
    # (labels-dir 모드는 main.py가 만들어둔 결과만 읽고 비교하므로
    # Gemini를 호출하지 않는다).
    # --------------------------------------------------------

    labels_dir: Optional[Path] = (
        Path(args.labels_dir) if args.labels_dir else None
    )

    if labels_dir is None and not GEMINI_API_KEY:

        raise SystemExit(
            "GEMINI_API_KEY가 설정되어 있지 않습니다.\n"
            "(사전 계산된 결과로만 평가하려면 --labels-dir을 사용하세요.)"
        )

    if labels_dir is not None and not labels_dir.exists():

        raise SystemExit(
            f"--labels-dir 디렉터리를 찾을 수 없습니다: {labels_dir}"
        )

    # --------------------------------------------------------
    # 경로
    # --------------------------------------------------------

    input_dir = Path(
        args.input_dir
    )

    profile_file = Path(
        args.profile_file
    )

    output_dir = Path(
        args.output_dir
    )

    output_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    # --------------------------------------------------------
    # profiles.json 로딩
    # --------------------------------------------------------

    profiles = load_profiles(
        profile_file
    )

    profile_map = {
        profile["profile_id"]: profile
        for profile in profiles
    }

    print(
        f"프로필 {len(profile_map)}개 로드 완료"
    )

    print(
        "프로필:",
        ", ".join(
            sorted(profile_map.keys())
        ),
    )

    # --------------------------------------------------------
    # 평가 대상 결정
    # --------------------------------------------------------

    if args.input:

        input_paths = [
            Path(args.input)
        ]

    else:

        if not input_dir.exists():

            raise SystemExit(
                f"입력 디렉터리를 찾을 수 없습니다: "
                f"{input_dir}"
            )

        # CSV + XLSX 모두 검색
        input_paths = sorted(
            [
                *input_dir.glob("P*.csv"),
                *input_dir.glob("P*.xlsx"),
            ],
            key=lambda path: path.stem,
        )

        if args.profile:

            input_paths = [
                path
                for path in input_paths
                if path.stem == args.profile
            ]

    # --------------------------------------------------------
    # CSV/XLSX 없음
    # --------------------------------------------------------

    if not input_paths:

        raise SystemExit(
            "평가할 CSV 또는 XLSX 파일을 찾지 못했습니다.\n"
            f"검색 위치: {input_dir}\n"
            "지원 파일: P*.csv, P*.xlsx"
        )

    # --------------------------------------------------------
    # 평가 대상 출력
    # --------------------------------------------------------

    print()
    print("평가 대상:")

    for path in input_paths:
        print(
            f"  - {path}"
        )

    # --------------------------------------------------------
    # Labeler
    #
    # labels-dir 모드에서는 Gemini를 호출하지 않으므로 생성하지 않는다.
    # --------------------------------------------------------

    labeler = (
        GroundTruthLabeler() if labels_dir is None else None
    )

    # 전체 결과
    all_rows: List[
        Dict[str, Any]
    ] = []

    # --------------------------------------------------------
    # 프로필별 평가
    # --------------------------------------------------------

    for input_path in input_paths:

        profile_name = input_path.stem

        # profiles.json에 없는 프로필
        if profile_name not in profile_map:

            print(
                f"[경고] {profile_name} 프로필이 "
                f"profiles.json에 없습니다. "
                f"건너뜁니다.",
                file=sys.stderr,
            )

            continue

        profile_data = profile_map[
            profile_name
        ]

        # ----------------------------------------------------
        # LLM에 전달할 profile_text
        # ----------------------------------------------------

        profile_text = str(
            profile_data.get(
                "profile_text",
                "",
            )
        ).strip()

        if not profile_text:

            print(
                f"[경고] {profile_name}의 "
                f"profile_text가 비어 있습니다. "
                f"건너뜁니다.",
                file=sys.stderr,
            )

            continue

        # ----------------------------------------------------
        # 평가
        # ----------------------------------------------------

        try:

            rows = await evaluate_profile(
                labeler=labeler,
                input_path=input_path,
                profile=profile_text,
                profile_name=profile_name,
                output_dir=output_dir,
                labels_dir=labels_dir,
            )

            all_rows.extend(
                rows
            )

        except Exception as e:

            print(
                f"[실패] {profile_name}: {e}",
                file=sys.stderr,
            )

    # --------------------------------------------------------
    # 결과 없음
    # --------------------------------------------------------

    if not all_rows:

        raise SystemExit(
            "평가 결과가 없습니다."
        )

    # --------------------------------------------------------
    # Overall
    # --------------------------------------------------------

    overall = calculate_metrics(
        all_rows
    )

    # --------------------------------------------------------
    # Per Profile
    # --------------------------------------------------------

    per_profile = calculate_per_profile(
        all_rows
    )

    # --------------------------------------------------------
    # 전체 불일치
    # --------------------------------------------------------

    disagreements = get_disagreements(
        all_rows
    )

    # --------------------------------------------------------
    # Overall 결과 저장
    # --------------------------------------------------------

    write_json(
        overall,
        output_dir
        / "overall_metrics.json",
    )

    # --------------------------------------------------------
    # Per Profile 결과 저장
    # --------------------------------------------------------

    write_json(
        per_profile,
        output_dir
        / "per_profile_metrics.json",
    )

    # --------------------------------------------------------
    # 전체 비교 결과
    # --------------------------------------------------------

    write_csv(
        all_rows,
        output_dir
        / "all_comparison.csv",
    )

    # --------------------------------------------------------
    # 전체 불일치
    # --------------------------------------------------------

    write_csv(
        disagreements,
        output_dir
        / "all_disagreements.csv",
    )

    # --------------------------------------------------------
    # 콘솔 출력
    # --------------------------------------------------------

    print_metrics(
        "OVERALL",
        overall,
    )

    print_per_profile(
        per_profile,
    )

    # --------------------------------------------------------
    # 완료
    # --------------------------------------------------------

    print()
    print(
        "=" * 60
    )

    print(
        "평가 완료"
    )

    print(
        "=" * 60
    )

    print(
        f"전체 비교 결과: "
        f"{output_dir / 'all_comparison.csv'}"
    )

    print(
        f"불일치 결과: "
        f"{output_dir / 'all_disagreements.csv'}"
    )

    print(
        f"Overall 결과: "
        f"{output_dir / 'overall_metrics.json'}"
    )

    print(
        f"Per-profile 결과: "
        f"{output_dir / 'per_profile_metrics.json'}"
    )


# ============================================================
# Entry Point
# ============================================================

def main() -> None:

    asyncio.run(
        async_main()
    )


if __name__ == "__main__":
    main()